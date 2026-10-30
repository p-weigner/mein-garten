"""Deterministischer EIVE-CC-BY-Offline-Export aus gepinntem Primärdatensatz.

GitHub Actions: einmalig Originalquelle prüfen, danach ein kleines Textpaket
erzeugen. NIEMALS OpenPlantbook-Freitexte, gärtnerische Pflegeaufgaben,
persönliche Notizen oder Home-Assistant-Zugangsdaten exportieren.
"""
from __future__ import annotations

import base64
import gzip
import hashlib
import io
import json
from pathlib import Path

from openpyxl import load_workbook

from tools.eive_europa_artenabdeckung import BENCHMARK, analyse_rows
from tools.eive_europa_quellenprobe import (
    DATA_URL, DOI, EXPECTED_MD5, EXPECTED_SIZE_MAX, METADATA_URL, ROOT,
    bounded_download, lizenz_id,
)

DEST = ROOT / "berichte" / "eive_europa_ccby_v1.json.gz.b64"
DIM = ("M", "N", "R", "L", "T")
ALLOWED_TAXON_SOURCES = {"EuroMed", "additionalTaxa"}
CATALOG_STATUS = "wissenschaftliche_standortindikatoren_ohne_pflegefreigabe"
LICENSE = "CC BY 4.0"


def produce_catalog(selected: dict) -> dict:
    """Nur konkrete Arten und belegte Einzelindikatoren, ohne Synonymersatz."""
    rows = []
    rejected = {}
    for key, item in sorted(selected.items()):
        according = item.get("taxon_herkunft")
        if according not in ALLOWED_TAXON_SOURCES:
            rejected[str(according)[:80]] = rejected.get(str(according)[:80], 0) + 1
            continue
        species = item["art"]
        cols = [species]
        for dim in DIM:
            f = item["felder"].get(dim)
            cols.extend([f["wert"], f["systeme"]] if f else [None, None])
        rows.append(cols)
    if not rows:
        raise ValueError("Keine nach Taxonquelle geeigneten EIVE-Arten")
    return {
        "schema": 1,
        "quelle": "EIVE 1.0",
        "lizenz": LICENSE,
        "doi": DOI,
        "herkunft": "https://zenodo.org/records/7534792",
        "original_md5": EXPECTED_MD5,
        "arten": rows,
        "felder": list(DIM),
        "pruefstatus": CATALOG_STATUS,
        "ausgeschlossene_arten_nach_taxonquelle": rejected,
    }


def serialize_catalog(catalog: dict) -> tuple[bytes, dict]:
    """Stabile UTF-8-JSON-Bytes, deterministische GZIP-Kompression, ASCII-Text."""
    body = (
        json.dumps(catalog, ensure_ascii=False, separators=(",", ":"), sort_keys=True)
        + "\n"
    ).encode("utf-8")
    compressed = gzip.compress(body, compresslevel=9, mtime=0)
    encoded = base64.b64encode(compressed)
    if len(encoded) >= 700_000 or len(body) >= 2_500_000:
        raise ValueError("EIVE-Katalog überschreitet das Zielgrößenlimit")
    return encoded + b"\n", {
        "catalog_uncompressed_bytes": len(body),
        "catalog_gzip_bytes": len(compressed),
        "catalog_base64_bytes": len(encoded),
        "catalog_utf8_sha256": hashlib.sha256(body).hexdigest(),
        "catalog_gzip_sha256": hashlib.sha256(compressed).hexdigest(),
        "catalog_base64_sha256": hashlib.sha256(encoded).hexdigest(),
        "catalog_arten": len(catalog["arten"]),
    }


def main():
    metadata = json.loads(bounded_download(METADATA_URL, 250_000))
    lizenz_id(metadata)
    source = bounded_download(DATA_URL, EXPECTED_SIZE_MAX)
    if hashlib.md5(source).hexdigest() != EXPECTED_MD5:
        raise ValueError("Unverifizierte EIVE-Originaldatei; Export abgebrochen")
    benchmark = json.loads(BENCHMARK.read_text(encoding="utf-8"))
    book = load_workbook(io.BytesIO(source), read_only=True, data_only=True)
    try:
        if "mainTable" not in book.sheetnames:
            raise ValueError("EIVE-Quelle enthält keine mainTable")
        sheet = book["mainTable"]
        sheet.reset_dimensions()
        measured = analyse_rows(sheet.iter_rows(values_only=True), benchmark)
    finally:
        book.close()
    catalog = produce_catalog(measured["technisch_lesbare_werte"])
    encoded, info = serialize_catalog(catalog)
    DEST.parent.mkdir(parents=True, exist_ok=True)
    DEST.write_bytes(encoded)
    bekannte_arten = {row[0].casefold() for row in catalog["arten"]}
    treffende_beispiele = sum(
        plant["wissenschaftlich"].casefold() in bekannte_arten for plant in benchmark
    )
    print("EIVE_EXPORT", json.dumps({
        **info, "benchmark50_exakt_vor_taxonpruefung": measured["benchmark50_exakt"],
        "benchmark50_exakt_nach_taxonpruefung": treffende_beispiele,
        "ausgeschlossene_quellen": catalog["ausgeschlossene_arten_nach_taxonquelle"],
        "lizenz": LICENSE, "doi": DOI, "automatische_pflegeaufgaben": 0,
    }, ensure_ascii=False), flush=True)
    # Zusätzliche Übergabe über das CI-Log: jedes Segment bleibt deutlich unter
    # der GitHub-Zeilenlänge. Nur CC-BY-Daten, keine privaten Informationen.
    encoded_text = encoded.decode("ascii").strip()
    segments = [encoded_text[i:i + 5000]
                for i in range(0, len(encoded_text), 5000)]
    print("EIVE_TRANSFER_BEGIN " + str(len(segments)), flush=True)
    for i, segment in enumerate(segments, 1):
        print("EIVE_TRANSFER_PART " + str(i) + " " + segment, flush=True)
    print("EIVE_TRANSFER_END " + info["catalog_base64_sha256"], flush=True)


if __name__ == "__main__":
    main()
