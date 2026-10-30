"""Botanischer Qualitäts- und Abdeckungstest der EIVE-1.0-Europadaten.

Keine Produktivintegration. Die Originaldatei wird ausschließlich im
GitHub-Runner geprüft. Weder unbearbeitete Rohdaten noch Pflegeanweisungen
werden in den Bericht übernommen.
"""
from __future__ import annotations

from collections import Counter
from hashlib import md5
import io
import json
import math
from pathlib import Path
import re

from tools.eive_europa_quellenprobe import (
    DATA_URL, DOI, EXPECTED_MD5, EXPECTED_SIZE_MAX, METADATA_URL,
    REQUIRED_LICENSE, ROOT, bounded_download, lizenz_id,
)

BASE = "EIVEres-"
DIMENSIONS = ("M", "N", "R", "L", "T")
REQUIRED = ("TaxonConcept", "TaxonRank", "AccordingTo") + tuple(
    x for dimension in DIMENSIONS
    for x in (BASE + dimension, BASE + dimension + ".n")
)
SCI_NAME = re.compile(r"^[A-Z][A-Za-z-]+ [a-z-]+$")
REPORT = ROOT / "berichte" / "eive_europa_artenabdeckung_20261030.json"
BENCHMARK = ROOT / "daten" / "pflanzensuche_50_de.json"


def number(value: object) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    if not math.isfinite(value) or not 0 <= value <= 10:
        return None
    return round(float(value), 2)


def analyse_rows(rows, benchmark: list[dict]) -> dict:
    """Liest alle Tabellenzeilen; nur exakt zugeordnete Art, ohne Synonymraten."""
    iterator = iter(rows)
    header = next(iterator, None)
    if not isinstance(header, (tuple, list)):
        raise ValueError("EIVE-Schema ohne Kopfzeile")
    if len(set(header)) != len(header):
        raise ValueError("EIVE-Schema mit doppelten Spalten")
    column = {key: header.index(key) for key in REQUIRED if key in header}
    if len(column) != len(REQUIRED):
        raise ValueError("EIVE-Schema: erforderliche Spalten fehlen")
    if any(column[key] >= 120 for key in REQUIRED):
        raise ValueError("EIVE-Schema hat unerwartet viele Spalten")

    selected = {}
    ambiguous = set()
    rejected = Counter()
    per_dimension = Counter()
    for row in iterator:
        if not isinstance(row, (tuple, list)):
            rejected["kein_datensatz"] += 1
            continue
        def cell(name):
            pos = column[name]
            return row[pos] if len(row) > pos else None

        name = cell("TaxonConcept")
        rank = cell("TaxonRank")
        if rank != "Species" or not isinstance(name, str) or not SCI_NAME.fullmatch(name):
            rejected["nicht_exakte_art"] += 1
            continue
        key = name.casefold()
        if key in ambiguous:
            rejected["art_mehrdeutig"] += 1
            continue
        if key in selected:
            ambiguous.add(key)
            del selected[key]
            rejected["art_mehrdeutig"] += 2
            continue

        values = {}
        for dimension in DIMENSIONS:
            value = number(cell(BASE + dimension))
            n = cell(BASE + dimension + ".n")
            if value is None or type(n) is not int or not 1 <= n <= 31:
                continue
            values[dimension] = {"wert": value, "systeme": n}
        if not values:
            rejected["keine_auswertbaren_indikatoren"] += 1
            continue
        selected[key] = {"art": name, "felder": values, "taxon_herkunft": cell("AccordingTo")}
        for dimension in values:
            per_dimension[dimension] += 1

    examples = []
    for plant in benchmark:
        key = plant["wissenschaftlich"].casefold()
        hit = selected.get(key)
        examples.append({
            "gruppe": plant["gruppe"],
            "name_de": plant["name_de"],
            "wissenschaftlich": plant["wissenschaftlich"],
            "status": "exakte_art" if hit else (
                "mehrdeutig" if key in ambiguous else "nicht_enthalten"
            ),
            "indikatoren": sorted(hit["felder"]) if hit else [],
            "quellsysteme_pro_indikator": {
                dim: value["systeme"]
                for dim, value in hit["felder"].items()
            } if hit else {},
        })
    return {
        "artzahl_eindeutig": len(selected),
        "arten_mehrdeutig": len(ambiguous),
        "verworfene_zeilen": dict(rejected),
        "arten_pro_indikator": dict(per_dimension),
        "benchmark50": examples,
        "benchmark50_exakt": sum(x["status"] == "exakte_art" for x in examples),
        "benchmark50_alle_fuenf_indikatoren": sum(
            len(x["indikatoren"]) == 5 for x in examples
        ),
        "technisch_lesbare_werte": selected,
    }


def main() -> None:
    from openpyxl import load_workbook

    metadata = json.loads(bounded_download(METADATA_URL, 250_000))
    lizenz_id(metadata)
    binary = bounded_download(DATA_URL, EXPECTED_SIZE_MAX)
    if md5(binary).hexdigest() != EXPECTED_MD5:
        raise ValueError("Originaldaten haben nicht die gepinnte EIVE-Prüfsumme")
    benchmark = json.loads(BENCHMARK.read_text(encoding="utf-8"))
    book = load_workbook(io.BytesIO(binary), read_only=True, data_only=True)
    try:
        if "mainTable" not in book.sheetnames:
            raise ValueError("EIVE-Tabelle mainTable fehlt")
        sheet = book["mainTable"]
        sheet.reset_dimensions()
        result = analyse_rows(sheet.iter_rows(values_only=True), benchmark)
    finally:
        book.close()

    report = {
        "source": "EIVE 1.0",
        "doi": DOI,
        "original_md5": EXPECTED_MD5,
        "license": REQUIRED_LICENSE,
        "status": "fachliche_einordnung_nicht_freigegeben",
        "hinweis": (
            "Europäische ökologische Nischenindikatoren sind keine "
            "kulturspezifischen Anweisungen für den Garten, keine Bewässerungs- "
            "oder Schnittintervalle und keine lokale Freiland-Prognose."
        ),
        "automatische_pflegeaufgaben": 0,
        **{key: value for key, value in result.items()
           if key != "technisch_lesbare_werte"},
    }
    REPORT.parent.mkdir(exist_ok=True, parents=True)
    REPORT.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print("EIVE_ABDECKUNG", json.dumps(
        {k: v for k, v in report.items() if k != "benchmark50"},
        ensure_ascii=False,
    ), flush=True)
    for row in report["benchmark50"]:
        print("EIVE_TESTART", json.dumps(row, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
